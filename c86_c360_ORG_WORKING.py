# =============================================================================
# C86 Client360 Product Appropriateness - SAS to Python Migration
# =============================================================================
# Description:
#   Full pipeline replicating the SAS-based C86 Client360 Product Appropriateness
#   workflow. Pulls data from Teradata, processes it in-memory with pandas,
#   and exports results to Excel and SQL Server.
#
# Steps:
#   1.  Setup Logging, Paths, and Configuration
#   2.  Helper Functions (Connections, Date Logic, Rationale Validation)
#   3.  Main Script Execution
#   4.  Pull tracking data (DDWV01.EVNT_PROD_TRACK_LOG)
#   5.  Process tracking data (in-memory)
#   6.  Pull C360 Detail Data (Volatile Table)
#   7.  Join Tool Info
#   8.  PROC FREQ replication
#   9.  Format ($Stagefmt) definition
#   10. Pull AOT Data
#   11. Process AOT Data (in-memory)
#   12. Create c360_detail_link_aot
#   13. Create c360_detail_more
#   14. PA Rationale Validation
#   15. Create C360_detail_more_in
#   16. Format ($cs_cmt) definition
#   17. Deduplicate data (tmp0)
#   18. Create base AC table (tmp_pa_C360_4ac)
#   19. Create AC Assessment (pa_C360_autocomplete_tool_use)
#   20. Prepare data for Tool Use Count
#   21. Create Tool Count Assessment
#   22. Combine AC datasets
#   23. (Commented out) Append and Save Autocomplete Data
#   24. Export Autocomplete to Excel
#   25. Create and Export Detail File
#   26. Export Pivot Table
#   SQL Load: Insert combined autocomplete data into SQL Server staging table
# =============================================================================

# --- Package Installation Check ---
import importlib
import subprocess
import sys

packages = ['teradatasql', 'pandas', 'numpy', 'openpyxl', 'pyarrow']

def install(package):
    """Installs a package using pip."""
    print(f"Installing {package}...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", package])

for package in packages:
    try:
        importlib.import_module(package)
        print(f"{package} is already installed.")
    except ImportError:
        print(f"{package} not found.")
        install(package)

print("\n--- Environment setup complete. You can now run the main script. ---")


# =============================================================================
# Imports
# =============================================================================
import pandas as pd
import numpy as np
import teradatasql
import json
import logging
import os
import re
import shutil
from datetime import datetime, date, timedelta
from pathlib import Path
from typing import Dict, List, Optional, Tuple


# =============================================================================
# ### 1. Setup Logging, Paths, and Configuration
# =============================================================================

def setup_logging(log_path: Path, log_file_name: str) -> Path:
    """Configures the logging module to write to a file and print to console."""
    log_path.mkdir(parents=True, exist_ok=True)
    log_file = log_path / log_file_name

    # Remove existing handlers to avoid duplicate logs in notebook re-runs
    for handler in logging.root.handlers[:]:
        logging.root.removeHandler(handler)

    logging.basicConfig(
        level=logging.INFO,
        format='%(asctime)s - %(levelname)s - %(message)s',
        handlers=[
            logging.FileHandler(log_file),
            logging.StreamHandler()  # This will print to the notebook output
        ]
    )
    logging.info(f"Logging configured. Log file: {log_file}")
    return log_file


# --- Configuration ---
# Per user request, hardcode the main output path
OUTPATH = Path("sas/RSD/REG_DEV/C86/output/product_appropriateness/client360")
LOGPATH = Path("sas/RSD/REG_DEV/C86/log/product_appropriateness/client360")

YMD = datetime.now().strftime('%Y%m%d')
LOG_FILE_NAME = f"C86_pa_client360_{YMD}.log"
LOG_FILE = setup_logging(LOGPATH, LOG_FILE_NAME)

logging.info(f"OUTPATH set to: {OUTPATH}")
logging.info(f"LOGPATH set to: {LOGPATH}")

# --- Create Output Directories ---
# Replicates %CreateDirectory(&outpath)
OUTPATH.mkdir(parents=True, exist_ok=True)
logging.info(f"Ensured output directory exists: {OUTPATH}")

# Define paths for 'libnames'
# libname ac "&outpath";
AC_PATH = OUTPATH
# libname dataout "&outpath/&runday";  (runday is defined later)


# =============================================================================
# ### 2. Helper Functions (Connections, Date Logic, Rationale)
# =============================================================================

def get_teradata_connection():
    """Reads connection details from JSON and returns a teradatasql connection."""
    conn_file = Path("TeradataConnection_T.json")
    if not conn_file.exists():
        logging.error(f"Connection file not found at: {conn_file.resolve()}")
        raise FileNotFoundError(f"Connection file not found at: {conn_file.resolve()}")

    try:
        with open(conn_file, 'r') as f:
            conn_details = json.load(f)

        logging.info(
            f"Connecting to Teradata URL: {conn_details.get('url')} "
            f"with user: {conn_details.get('user')} using LDAP..."
        )
        conn = teradatasql.connect(
            host=conn_details.get('url'),
            user=conn_details.get('user'),
            password=conn_details.get('password'),
            logmech="LDAP"
        )
        logging.info("Teradata connection successful.")
        return conn
    except Exception as e:
        logging.error(f"Failed to connect to Teradata: {e}")
        raise


def check_ini_run(outpath, autocomplete_filename) -> str:
    """
    Checks if the persistent autocomplete file exists to determine ini_run status.
    Replicates the %ini_check macro logic.
    """
    autocomplete_file = Path(outpath) / autocomplete_filename
    autocomplete_backup = Path(outpath) / f"{autocomplete_file.stem}_backup{autocomplete_file.suffix}"

    if autocomplete_file.exists():
        logging.info(f"Found existing autocomplete file: {autocomplete_file}")
        try:
            shutil.copy(autocomplete_file, autocomplete_backup)
            logging.info(f"Backed up to: {autocomplete_backup}")
        except Exception as e:
            logging.warning(f"Could not create backup of autocomplete file: {e}")
        return 'N'
    else:
        logging.info("Autocomplete file not found. Setting ini_run = 'Y'.")
        return 'Y'


def setup_dates(ini_run: str) -> Tuple[Dict[str, str], date]:
    """
    Replicates the main _null_ data step to set up all date macro variables.
    """
    tday = date.today()
    runday = tday.strftime('%Y%m%d')

    # SAS: intnx('week.4', tday, 0, 'b')
    # week.4 = weeks start on Wednesday (weekday() == 2)
    # 'b' = beginning
    days_to_wednesday = (tday.weekday() - 2) % 7
    intnx_week4_b = tday - timedelta(days=days_to_wednesday)

    week_start_dt = intnx_week4_b - timedelta(days=12)
    week_end_dt = intnx_week4_b - timedelta(days=6)

    dates: Dict[str, str] = {}

    if ini_run == 'Y':
        logging.info("ini_run = 'Y'. Using initial launch dates.")
        launch_dt = date(2023, 5, 7)
        launch_dt_min14 = date(2023, 4, 23)

        dates['wk_start'] = f"'{launch_dt.strftime('%Y-%m-%d')}'"
        dates['wk_start_min14'] = f"'{launch_dt_min14.strftime('%Y-%m-%d')}'"
        dates['wk_end'] = f"'{week_end_dt.strftime('%Y-%m-%d')}'"
    else:  # 'N' or any other value
        logging.info("ini_run = 'N'. Using rolling weekly dates.")
        dates['wk_start'] = f"'{week_start_dt.strftime('%Y-%m-%d')}'"
        dates['wk_start_min14'] = f"'{(week_start_dt - timedelta(days=14)).strftime('%Y-%m-%d')}'"
        dates['wk_end'] = f"'{week_end_dt.strftime('%Y-%m-%d')}'"

    dates['runday'] = runday
    dates['tday'] = runday

    logging.info(f"Dates calculated: {dates}")
    return dates, week_end_dt


def validate_rationale(text):
    """
    Replicates the pa_rationale data step logic to validate text.
    Returns: (category, xfail_chars_gt5, xfail_rep_char, xfail_ge_2_alnum)
    """
    if pd.isna(text) or not isinstance(text, str):
        return "Invalid", 1, 1, 1

    # SAS: _x = upcase(strip(translate(_x,' ','_x_p')));
    # This normalizes whitespace, strips, and uppercases.
    _x = ' '.join(text.split()).upper()

    # (1) xfail_chars_gt5: > 5 characters
    xfail_chars_gt5 = 0 if len(_x) > 5 else 1

    # (2) xfail_rep_char: not only repeated characters
    # SAS: _x2 = compress(_x, substrn(_x, 1, 1));
    _x2 = ""
    if len(_x) > 0:
        _x2 = _x.replace(_x[0], '')
    # SAS: xfail_rep_char = ifn(not missing(_x2), 0, 1);
    xfail_rep_char = 0 if len(_x2) > 0 else 1

    # (3) xfail_ge_2_alnum: have at least 2 alphabet/num characters
    # SAS: _x3 = compress(_x, 'a', 'kad');
    _x3 = re.sub(r'[^a-zA-Z0-9]', '', _x)
    # SAS: xfail_ge_2_alnum = ifn(lengthn(_x3) > 2, 0, 1);
    # NOTE: The SAS code says > 2. This seems like a bug in the SAS.
    # "at least 2" should be "len(_x3) >= 2", which SAS would be "lengthn(_x3) > 1".
    # Replicating the SAS logic (lengthn(_x3) > 2) exactly.
    xfail_ge_2_alnum = 0 if len(_x3) > 2 else 1

    # Final category
    # SAS: prod_not_aprp_rtnl_txt_cat = ifc(sum(of xfail_:) = 0, "Valid", "Invalid");
    is_valid = (xfail_chars_gt5 + xfail_rep_char + xfail_ge_2_alnum) == 0
    category = "Valid" if is_valid else "Invalid"

    return category, xfail_chars_gt5, xfail_rep_char, xfail_ge_2_alnum


# =============================================================================
# ### 3. Main Script Execution
# =============================================================================

def main() -> None:

    ymd = datetime.now().strftime("%Y%m%d")
    log_file_name = f"C86_pa_client360_{ymd}.log"
    setup_logging(LOGPATH, log_file_name)

    # Using a persistent file as the Python equivalent of the SAS dataset
    PERSISTENT_AC_FILE = "pa_client360_autocomplete.parquet"

    try:
        logging.info("=" * 50)
        logging.info("SAS to Python Migration: C86_pa_client360 Starting...")
        logging.info(f"Script running as user: {os.environ.get('USER', 'Unknown')}")
        logging.info(f"Platform: {os.environ.get('HOSTNAME', 'Unknown')}")

        # --- Inicheck logic ---
        # Replicates %ini_check macro
        # We run the logic but then obey the user request to force 'N' for testing.
        ini_run_from_logic = check_ini_run(AC_PATH, PERSISTENT_AC_FILE)

        # PER USER REQUEST: Force 'N' for testing, but show the logic was included.
        ini_run = 'N'
        logging.info(
            f"ini_run logic determined: '{ini_run_from_logic}'. "
            f"FORCING '{ini_run}' for testing."
        )

        # --- Date setup ---
        # Replicates the _null_ data step
        dates, week_end_dt = setup_dates(ini_run)
        runday = dates["runday"]

        # --- Create runday directory ---
        # Replicates %CreateDirectory(&outpath/&runday)
        DATAOUT_PATH = OUTPATH / runday
        DATAOUT_PATH.mkdir(parents=True, exist_ok=True)
        logging.info(f"Ensured runday output directory exists: {DATAOUT_PATH}")

        # -------------------------------------------------------------------------
        # --- 4. Pull tracking data ---
        # -------------------------------------------------------------------------
        logging.info("Step 4: Pulling data from tracking_all (DDWV01.EVNT_PROD_TRACK_LOG)...")

        # Note: SAS `date &WK_START - 90` becomes
        # `CAST({dates['wk_start']} AS DATE) - 90` in Teradata SQL
        sql_tracking_all = f"""
SELECT *
FROM DDWV01.EVNT_PROD_TRACK_LOG
WHERE advr_selt_typ = 'Advice Tool'
  AND EVNT_DT > (CAST({dates['wk_start']} AS DATE) - 90)
"""
        with get_teradata_connection() as conn:
            tracking_all = pd.read_sql(sql_tracking_all, conn)

        logging.info(f"Loaded {len(tracking_all)} rows into tracking_all.")

        # -------------------------------------------------------------------------
        # --- 5. Process tracking data (in-memory) ---
        # -------------------------------------------------------------------------
        logging.info("Step 5: Processing tracking data (replicating 3 PROC SQL steps)...")

        # Filter once for all subsequent steps
        # SAS: where OPPOR_ID is not missing and ADVC_TOOL_NM is not missing
        tracking_all_filtered = tracking_all.dropna(
            subset=['OPPOR_ID', 'ADVC_TOOL_NM']
        ).copy()

        # 1. tracking_tool_use_distinct
        # SAS: select distinct OPPOR_ID, upcase(ADVC_TOOL_NM) as ADVC_TOOL_NM
        tracking_all_filtered['ADVC_TOOL_NM'] = tracking_all_filtered['ADVC_TOOL_NM'].str.upper()
        tracking_tool_use_distinct = (
            tracking_all_filtered[['OPPOR_ID', 'ADVC_TOOL_NM']].drop_duplicates()
        )
        logging.info(f"Created tracking_tool_use_distinct: {len(tracking_tool_use_distinct)} rows")

        # 2. tracking_count_tool_use_pre2
        # SAS: select OPPOR_ID, count(distinct upcase(ADVC_TOOL_NM)) as count_unique_tool_used
        tracking_count_tool_use_pre2 = (
            tracking_all_filtered
            .groupby('OPPOR_ID')['ADVC_TOOL_NM']
            .nunique()
            .reset_index(name='count_unique_tool_used')
            .sort_values(by='count_unique_tool_used', ascending=False)
        )
        logging.info(f"Created tracking_count_tool_use_pre2: {len(tracking_count_tool_use_pre2)} rows")

        # 3. tracking_tool_use
        # SAS: case when count_unique_tool_used > 0 then 'Tool Used' end as tool_used
        tracking_tool_use = tracking_count_tool_use_pre2[
            tracking_count_tool_use_pre2['count_unique_tool_used'] > 0
        ][['OPPOR_ID']].copy()
        tracking_tool_use['tool_used'] = 'Tool Used'
        logging.info(f"Created tracking_tool_use: {len(tracking_tool_use)} rows")

        # -------------------------------------------------------------------------
        # --- 6. Pull C360 Detail Data (Volatile Table) ---
        # -------------------------------------------------------------------------
        logging.info("Step 6: Pulling C360 detail data (using Volatile Table)...")

        # This process must be in a single session
        sql_create_vt_c360_short = f"""
CREATE MULTISET VOLATILE TABLE c360_short AS
(
  SELECT evnt_id,
         CAST(rbc_oppor_own_id AS INTEGER) AS emp_id,
         evnt_dt AS snap_dt
  FROM ddwv01.evnt_prod_oppor
  WHERE rbc_oppor_own_id IS NOT NULL
    AND evnt_dt IS NOT NULL
    AND evnt_id IS NOT NULL
    AND evnt_dt BETWEEN {dates['wk_start']} AND {dates['wk_end']}
)
WITH DATA PRIMARY INDEX (emp_id, snap_dt) ON COMMIT PRESERVE ROWS;
"""
        sql_stats_c360_short = (
            "COLLECT STATISTICS COLUMN (emp_id, snap_dt) ON c360_short;"
        )

        sql_select_c360_detail_pre = f"""
SELECT c360.*,
       emp.org_unt_no, emp.hr_posn_titl_en,
       emp.posn_strt_dt, emp.posn_end_dt, emp.occpt_job_cd
FROM ddwv01.evnt_prod_oppor AS c360
LEFT JOIN
(
  SELECT c3.evnt_id,
         e1.org_unt_no, e1.hr_posn_titl_en,
         e2.posn_strt_dt, e2.posn_end_dt, e1.occpt_job_cd
  FROM c360_short AS c3
  INNER JOIN ddwv01.emp AS e1
    ON  e1.emp_id   = c3.emp_id
    AND c3.snap_dt >= e1.captr_dt
    AND c3.snap_dt  < e1.chg_dt
  INNER JOIN ddwv01.empl_reltn AS e2
    ON  e2.emp_id   = c3.emp_id
    AND c3.snap_dt >= e2.captr_dt
    AND c3.snap_dt  < e2.chg_dt
) AS emp
  ON emp.evnt_id = c360.evnt_id
WHERE c360.evnt_id IS NOT NULL
  AND c360.evnt_dt BETWEEN {dates['wk_start']} AND {dates['wk_end']};
"""
        with get_teradata_connection() as conn:
            with conn.cursor() as cursor:
                logging.info("Executing: CREATE VOLATILE TABLE c360_short...")
                cursor.execute(sql_create_vt_c360_short)
                logging.info("Executing: COLLECT STATISTICS...")
                cursor.execute(sql_stats_c360_short)
                logging.info("Executing: SELECT c360_detail_pre...")
                c360_detail_pre = pd.read_sql(sql_select_c360_detail_pre, conn)
                logging.info(f"Loaded {len(c360_detail_pre)} rows into c360_detail_pre.")
        # Volatile table is dropped when connection closes

        # -------------------------------------------------------------------------
        # --- 7. Join Tool Info ---
        # -------------------------------------------------------------------------
        logging.info("Step 7: Joining tool info to c360_detail_pre...")

        c360_detail = pd.merge(c360_detail_pre, tracking_tool_use, on='OPPOR_ID', how='left')

        # Replicate SAS case statement
        c360_detail['TOOL_USED'] = c360_detail['tool_used'].fillna('Tool Not Used')
        c360_detail = c360_detail.drop(columns=['tool_used'])  # Drop the original column

        logging.info(f"Created c360_detail: {len(c360_detail)} rows")

        if 'LOB' not in c360_detail.columns:
            c360_detail['LOB'] = 'Retail'
        else:
            c360_detail['LOB'] = c360_detail['LOB'].fillna('Retail')

        logging.info(f"Created c360_detail: {len(c360_detail)} rows with LOB filled as 'Retail'")

        # -------------------------------------------------------------------------
        # --- 8. PROC FREQ replication ---
        # -------------------------------------------------------------------------
        logging.info("PROC FREQ for c360_detail['LOB']:")
        logging.info(f"\n{c360_detail['LOB'].value_counts(dropna=False)}")

        # -------------------------------------------------------------------------
        # --- 9. Format ($Stagefmt) definition ---
        # -------------------------------------------------------------------------
        logging.info("Step 9: Defining $Stagefmt map...")
        stage_format_map = {
            "Démarche exploratoire/Comprendre le besoin": "11.Démarche exploratoire/Comprendre le besoin",
            "Discovery/Understand Needs":                 "12.Discovery/Understand Needs",
            "Review Options":                             "21.Review Options",
            "Present/Gain Commitment":                    "31.Present/Gain Commitment",
            "Intégration commencée":                      "41.Intégration commencée",
            "Onboarding Started":                         "42.Onboarding Started",
            "Opportunity Lost":                           "51.Opportunity Lost",
            "Opportunity Won":                            "61.Opportunity Won",
        }

        # -------------------------------------------------------------------------
        # --- 10. Pull AOT Data ---
        # -------------------------------------------------------------------------
        logging.info("Step 10: Pulling AOT data (ddwv01.evnt_prod_aot)...")
        sql_aot_all_oppor = f"""
SELECT
  oppor_id,
  COUNT(*) AS count_aot
FROM ddwv01.evnt_prod_aot
WHERE ess_src_evnt_dt BETWEEN {dates['wk_start_min14']} AND {dates['wk_end']}
  AND oppor_id IS NOT NULL
GROUP BY 1
"""
        with get_teradata_connection() as conn:
            aot_all_oppor = pd.read_sql(sql_aot_all_oppor, conn)

        logging.info(f"Loaded {len(aot_all_oppor)} rows into aot_all_oppor.")

        # -------------------------------------------------------------------------
        # --- 11. Process AOT Data (in-memory) ---
        # -------------------------------------------------------------------------
        aot_all_oppor_unique = aot_all_oppor[['OPPOR_ID']].drop_duplicates().copy()
        aot_all_oppor_unique.rename(columns={'OPPOR_ID': 'aot_oppor_id'}, inplace=True)
        logging.info(
            f"Created aot_all_oppor_unique: {len(aot_all_oppor_unique)} rows "
            f"with column 'aot_oppor_id'."
        )

        # -------------------------------------------------------------------------
        # --- 12. Create c360_detail_link_aot ---
        # -------------------------------------------------------------------------
        logging.info("Step 12: Creating c360_detail_link_aot...")

        c360_detail_link_aot = pd.merge(
            c360_detail,
            aot_all_oppor_unique,
            left_on='OPPOR_ID',
            right_on='aot_oppor_id',
            how='left'
        )

        cond_prod = c360_detail_link_aot['PROD_CATG_NM'] == 'Personal Accounts'
        cond_aot  = c360_detail_link_aot['aot_oppor_id'].notna()
        c360_detail_link_aot['C360_PDA_Link_AOT'] = np.where(cond_prod & cond_aot, 1, 0)

        logging.info(f"Created c360_detail_link_aot: {len(c360_detail_link_aot)} rows.")

        # -------------------------------------------------------------------------
        # --- 13. Create c360_detail_more ---
        # -------------------------------------------------------------------------
        logging.info("Step 13: Filtering to c360_detail_more_in_pre...")

        df = c360_detail_link_aot.copy()

        # Apply format
        df['oppor_stage_nm_f'] = df['OPPOR_STAGE_NM'].map(stage_format_map)

        # Define filter conditions for the 'in_pre' subset
        cond1 = df['ASCT_PROD_FMLY_NM'] != 'Risk Protection'
        cond2 = df['LOB'] == 'Retail'
        cond3 = df['C360_PDA_Link_AOT'] == 0
        cond4 = df['OPPOR_STAGE_NM'].isin(['Opportunity Won', 'Opportunity Lost'])

        all_conditions = cond1 & cond2 & cond3 & cond4
        c360_detail_more_in_pre = df[all_conditions].copy()

        # data c360_detail_more_i0; set c360_detail_more_in_pre;
        c360_detail_more_i0 = c360_detail_more_in_pre.copy()
        logging.info(f"Created c360_detail_more_i0: {len(c360_detail_more_i0)} rows.")

        # -------------------------------------------------------------------------
        # --- 14. PA Rationale Validation ---
        # -------------------------------------------------------------------------
        logging.info("Step 14: Running PA Rationale validation...")

        # Filter to relevant rows
        pa_rationale_base = c360_detail_more_i0[
            c360_detail_more_i0['IS_PROD_APRP_FOR_CLNT'] == 'Not Appropriate - Rationale'
        ].copy()

        pa_rationale = pa_rationale_base[['EVNT_ID', 'IS_PROD_APRP_FOR_CLNT', 'CLNT_RTNL_TXT']]

        # Apply validation function
        validation_results = pa_rationale['CLNT_RTNL_TXT'].apply(
            lambda x: pd.Series(
                validate_rationale(x),
                index=[
                    'prod_not_aprp_rtnl_txt_cat',
                    'xfail_chars_gt5',
                    'xfail_rep_char',
                    'xfail_ge_2_alnum'
                ]
            )
        )

        pa_rationale = pd.concat([pa_rationale, validation_results], axis=1)
        logging.info(f"Validated {len(pa_rationale)} rationale entries.")

        # -------------------------------------------------------------------------
        # --- 15. Create C360_detail_more_in ---
        # -------------------------------------------------------------------------
        logging.info("Step 15: Creating C360_detail_more_in...")

        # Join validation results back
        c360_detail_more_in = pd.merge(
            c360_detail_more_i0,
            pa_rationale.drop(columns=['IS_PROD_APRP_FOR_CLNT', 'CLNT_RTNL_TXT']),  # Avoid duplicate cols
            on='EVNT_ID',
            how='left'
        )

        # Replicate the final CASE statement
        conds = [
            c360_detail_more_in['IS_PROD_APRP_FOR_CLNT'].isna(),
            c360_detail_more_in['IS_PROD_APRP_FOR_CLNT'] != 'Not Appropriate - Rationale'
        ]
        choices = ['Not Available', 'Not Applicable']

        # Default is the calculated category from the join
        c360_detail_more_in['prod_not_aprp_rtnl_txt_cat'] = np.select(
            conds,
            choices,
            default=c360_detail_more_in['prod_not_aprp_rtnl_txt_cat']
        )

        logging.info(f"Created C360_detail_more_in: {len(c360_detail_more_in)} rows.")

        # -------------------------------------------------------------------------
        # --- 16. Format ($cs_cmt) definition ---
        # -------------------------------------------------------------------------
        logging.info("Step 16: Defining $cs_cmt map...")
        cs_cmt_map = {
            'COM1':  'Test population (less samples)',
            'COM2':  'Match population',
            'COM3':  'Mismatch population (less samples)',
            'COM4':  'Non Anomaly Population',
            'COM5':  'Anomaly Population',
            'COM6':  'Number of Deposit Sessions',
            'COM7':  'Number of Accounts',
            'COM8':  'Number of Transactions',
            'COM9':  'Non Blank Population',
            'COM10': 'Blank Population',
            'COM11': 'Unable to Assess',
            'COM12': 'Number of Failed Data Elements',
            'COM13': 'Population Distribution',
            'COM14': 'Reconciled Population',
            'COM15': 'Not Reconciled Population',
            'COM16': 'Pass',
            'COM17': 'Fail',
            'COM18': 'Not Applicable',
            'COM19': 'Potential Fail',
        }

        # -------------------------------------------------------------------------
        # --- 17. Deduplicate data (tmp0) ---
        # -------------------------------------------------------------------------
        logging.info("Step 17: Deduplicating by OPPOR_ID...")

        # proc sort + data step with first.
        tmp0 = c360_detail_more_in.sort_values(by='OPPOR_ID').copy()
        tmp0['level_oppor'] = tmp0.groupby('OPPOR_ID').cumcount() + 1

        # -------------------------------------------------------------------------
        # --- 18. Create base AC table (tmp_pa_C360_4ac) ---
        # -------------------------------------------------------------------------
        logging.info("Step 18: Creating base AC table tmp_pa_C360_4ac...")

        tmp_pa_c360_4ac = tmp0[tmp0['level_oppor'] == 1].copy()

        # Ensure EVNT_DT is datetime for calculations
        tmp_pa_c360_4ac['EVNT_DT'] = pd.to_datetime(tmp_pa_c360_4ac['EVNT_DT'])

        # --- START FIX ---
        # This logic correctly finds the "end of week" Friday (weekday=4)
        # for each row's EVNT_DT.
        logging.info("Calculating row-by-row SnapDate for end-of-week Friday...")

        # 1. Get the current weekday (Monday=0, Friday=4, Sunday=6)
        day_of_week = tmp_pa_c360_4ac['EVNT_DT'].dt.weekday

        # 2. Calculate days to add to get to Friday (4)
        # If today is Wed (2), 4 - 2 = 2 days to add.
        # If today is Fri (4), 4 - 4 = 0 days to add.
        # If today is Sat (5), 4 - 5 = -1. 'where' clause triggers: 7 + (-1) = 6 days
        #   to add (rolls to *next* Friday).
        days_to_friday = 4 - day_of_week
        days_to_friday = days_to_friday.where(days_to_friday >= 0, 7 + days_to_friday)

        # 3. Add the calculated days to the event date
        SnapDate_series = tmp_pa_c360_4ac['EVNT_DT'] + pd.to_timedelta(days_to_friday, unit='d')
        # --- END FIX ---

        # Replicate the large data step assignment
        tmp_pa_c360_4ac = tmp_pa_c360_4ac.assign(
            RegulatoryName='C86',
            LOB='Retail',
            ReportName='C86 Client360 Product Appropriateness',
            ControlRisk='Completeness',
            TestType='Anomaly',
            TestPeriod='Origination',
            ProductType=tmp_pa_c360_4ac['PROD_CATG_NM'],
            RDE='',               # Placeholder; overwritten per assessment in steps 19 & 21
            segment='Account Open',
            segment2=tmp_pa_c360_4ac['ASCT_PROD_FMLY_NM'],
            segment3=tmp_pa_c360_4ac['PROD_SRVC_NM'],
            segment4='',          # Placeholder; mapped from IS_PROD_APRP_FOR_CLNT in steps 19 & 21
            segment5='',          # Placeholder; mapped from prod_not_aprp_rtnl_txt_cat in steps 19 & 21
            segment6=tmp_pa_c360_4ac['OPPOR_STAGE_NM'],
            segment7=tmp_pa_c360_4ac['TOOL_USED'],
            segment8='',          # Placeholder; populated from ADVC_TOOL_NM in step 20
            segment9='',          # Placeholder; set explicitly in steps 19 & 21
            segment10=tmp_pa_c360_4ac['EVNT_DT'].dt.strftime('%Y%m'),
            CommentCode='COM13',
            Comments=cs_cmt_map['COM13'],
            HoldoutFlag='N',
            # Use the new SnapDate series; .dt.date converts it from datetime to just the date part.
            SnapDate=SnapDate_series.dt.date,
            DateCompleted=pd.to_datetime(dates['tday'], format='%Y%m%d').date()
        )

        logging.info(f"Created tmp_pa_C360_4ac: {len(tmp_pa_c360_4ac)} rows.")

        # -------------------------------------------------------------------------
        # --- 19. Create AC Assessment (pa_C360_autocomplete_tool_use) ---
        # -------------------------------------------------------------------------
        logging.info("Step 19: Creating AC assessment tmp_pa_C360_ac_assessment...")

        df_agg = tmp_pa_c360_4ac.copy()

        # segment4 mapping
        seg4_map = {
            'Product Appropriateness assessed outside Client 360': 'Product Appropriate assessed outside Client 360',
            'Not Appropriate - Rationale':                          'Product Not Appropriate',
            'Client declined product appropriateness assessment':   'Client declined product appropriateness assessment',
            'Product Appropriate':                                  'Product Appropriate',
        }

        df_agg['segment4'] = df_agg['IS_PROD_APRP_FOR_CLNT'].map(seg4_map).fillna('Missing')
        df_agg['segment5'] = df_agg['prod_not_aprp_rtnl_txt_cat']
        df_agg['RDE']      = 'PA002_Client360_Completeness_RDE'
        df_agg['segment9'] = ''

        print(df_agg.columns)

        group_by_cols = [
            'RegulatoryName', 'LOB', 'ReportName', 'ControlRisk', 'TestType',
            'TestPeriod', 'ProductType', 'RDE', 'segment', 'segment2',
            'segment3', 'segment4', 'segment5', 'segment6', 'segment7', 'segment8', 'segment9',
            'segment10', 'HoldoutFlag', 'CommentCode', 'Comments',
            'DateCompleted', 'SnapDate'
        ]

        df_agg['segment3'] = df_agg['segment3'].fillna("")

        # Fill NaN in group columns to match SAS grouping behavior
        # segment5 is left as-is (prod_not_aprp_rtnl_txt_cat can be NaN/blank legitimately)
        fill_cols = [c for c in group_by_cols if c != 'segment5']
        df_agg[fill_cols] = df_agg[fill_cols].fillna('Missing')

        tmp_pa_c360_ac_assessment = (
            df_agg
            .groupby(group_by_cols, dropna=False)
            .size()
            .reset_index(name='Volume')
        )

        tmp_pa_c360_ac_assessment['Amount'] = np.nan  # Replicate sum(.) as Amount

        tmp_pa_c360_ac_assessment['DateCompleted'] = pd.to_datetime(
            tmp_pa_c360_ac_assessment['DateCompleted'], errors='coerce'
        )
        tmp_pa_c360_ac_assessment['SnapDate'] = pd.to_datetime(
            tmp_pa_c360_ac_assessment['SnapDate'], errors='coerce'
        )

        # data pa_C360_autocomplete_tool_use; set ...;
        pa_c360_autocomplete_tool_use = tmp_pa_c360_ac_assessment.copy()

        logging.info(f"Created pa_c360_autocomplete_tool_use: {len(pa_c360_autocomplete_tool_use)} rows.")

        # -------------------------------------------------------------------------
        # --- 20. Prepare data for Tool Use Count ---
        # -------------------------------------------------------------------------
        logging.info("Step 20: Preparing data for tool use count...")

        # tmp_pa_C360_4ac_count_pre
        # Merge with tracking_tool_use_distinct to bring in ADVC_TOOL_NM per OPPOR_ID
        tmp_pa_c360_4ac_count_pre = pd.merge(
            tmp_pa_c360_4ac,
            tracking_tool_use_distinct,   # Created in Step 5: OPPOR_ID + ADVC_TOOL_NM
            on='OPPOR_ID',
            how='left',
            suffixes=('', '_tool')        # Avoid collision if ADVC_TOOL_NM already exists
        )

        # Resolve ADVC_TOOL_NM: prefer the joined value; fall back to any pre-existing column
        if 'ADVC_TOOL_NM_tool' in tmp_pa_c360_4ac_count_pre.columns:
            tmp_pa_c360_4ac_count_pre['ADVC_TOOL_NM'] = (
                tmp_pa_c360_4ac_count_pre['ADVC_TOOL_NM_tool']
                .fillna(tmp_pa_c360_4ac_count_pre.get('ADVC_TOOL_NM', pd.Series(dtype=str)))
            )
            tmp_pa_c360_4ac_count_pre.drop(columns=['ADVC_TOOL_NM_tool'], inplace=True)

        # tmp_pa_C360_4ac_count
        tmp_pa_c360_4ac_count = tmp_pa_c360_4ac_count_pre.copy()
        # segment8 = individual tool name (ADVC_TOOL_NM); blank when not available
        tmp_pa_c360_4ac_count['segment8'] = tmp_pa_c360_4ac_count['ADVC_TOOL_NM'].fillna('')
        # Note: segment10 format change is ignored as it's just metadata

        logging.info(f"Created tmp_pa_C360_4ac_count: {len(tmp_pa_c360_4ac_count)} rows.")

        # -------------------------------------------------------------------------
        # --- 21. Create Tool Count Assessment ---
        # -------------------------------------------------------------------------
        logging.info("Step 21: Creating tool count assessment...")

        df_agg_count = tmp_pa_c360_4ac_count.copy()

        # segment4 mapping (different from step 19)
        seg4_map_count = {
            'Not Appropriate - Rationale':                          'Product Not Appropriate',
            'Client declined product appropriateness assessment':   'Client declined',
            'Product Appropriate':                                  'Product Appropriate',
            'Product Appropriateness assessed outside Client 360':  'Product Appropriateness assessed outside Client 360',
        }

        df_agg_count['segment9'] = ''
        df_agg_count['segment4'] = df_agg_count['IS_PROD_APRP_FOR_CLNT'].map(seg4_map_count).fillna('Missing')
        df_agg_count['RDE']      = 'PA003_Client360_Completeness_Tool'

        # segment5 fix
        df_agg_count['segment5'] = df_agg_count['prod_not_aprp_rtnl_txt_cat']

        # Define the final grouping list one time
        group_by_cols_count = [
            'RegulatoryName', 'LOB', 'ReportName', 'ControlRisk', 'TestType', 'TestPeriod',
            'ProductType', 'RDE', 'segment', 'segment2', 'segment3',
            'segment4', 'segment5', 'segment6', 'segment7', 'segment8', 'segment9',
            'segment10', 'HoldoutFlag', 'CommentCode', 'Comments',
            'DateCompleted', 'SnapDate'
        ]

        # Fill NaNs one time
        # segment9 is left as empty string '' — excluded from fillna to preserve blank
        fill_cols_count = [c for c in group_by_cols_count if c != 'segment9']
        df_agg_count[fill_cols_count] = df_agg_count[fill_cols_count].fillna('Missing')

        # Group one time
        tmp_pa_c360_ac_count_assessment = (
            df_agg_count
            .groupby(group_by_cols_count, dropna=False)
            .size()
            .reset_index(name='Volume')
        )

        tmp_pa_c360_ac_count_assessment['Amount'] = np.nan

        tmp_pa_c360_ac_count_assessment['DateCompleted'] = pd.to_datetime(
            tmp_pa_c360_ac_count_assessment['DateCompleted'], errors='coerce'
        )
        tmp_pa_c360_ac_count_assessment['SnapDate'] = pd.to_datetime(
            tmp_pa_c360_ac_count_assessment['SnapDate'], errors='coerce'
        )

        # data pa_C360_autocomplete_Count_Tool; set ...;
        pa_c360_autocomplete_count_tool = tmp_pa_c360_ac_count_assessment.copy()

        logging.info(f"Created pa_C360_autocomplete_Count_Tool: {len(pa_c360_autocomplete_count_tool)} rows.")

        # -------------------------------------------------------------------------
        # --- 22. Combine AC datasets ---
        # -------------------------------------------------------------------------
        logging.info("Step 22: Combining autocomplete datasets...")

        for _df in (pa_c360_autocomplete_tool_use, pa_c360_autocomplete_count_tool):
            if 'segment9' not in _df.columns:
                _df['segment9'] = ''

        combine_pa_autocomplete = pd.concat(
            [pa_c360_autocomplete_count_tool, pa_c360_autocomplete_tool_use],
            ignore_index=True
        )

        combine_pa_autocomplete['DateCompleted'] = pd.to_datetime(
            combine_pa_autocomplete['DateCompleted'], errors='coerce'
        )
        combine_pa_autocomplete['SnapDate'] = pd.to_datetime(
            combine_pa_autocomplete['SnapDate'], errors='coerce'
        )

        logging.info(f"Created combine_pa_autocomplete: {len(combine_pa_autocomplete)} rows.")

        ac_pa_client360_autocomplete = combine_pa_autocomplete.copy()

        # -------------------------------------------------------------------------
        # --- 23. (Commented out) Append and Save Autocomplete Data ---
        # -------------------------------------------------------------------------
        # autocomplete_path = AC_PATH / PERSISTENT_AC_FILE
        # today_date = pd.to_datetime(dates['tday'], format='%Y%m%d').date()
        # if autocomplete_path.exists():
        #     logging.info(f"Appending to existing file: {autocomplete_path}")
        #     ac_pa_client360_autocomplete_old = pd.read_parquet(autocomplete_path)
        #     # Ensure DateCompleted is date object for comparison
        #     ac_pa_client360_autocomplete_old['DateCompleted'] = (
        #         pd.to_datetime(ac_pa_client360_autocomplete_old['DateCompleted']).dt.date
        #     )
        #     # SAS: where=(DateCompleted=input("&runday", anydtdte8.))
        #     # This is a *filter*, so we take everything NOT from today
        #     ac_pa_client360_autocomplete_filtered = ac_pa_client360_autocomplete_old[
        #         ac_pa_client360_autocomplete_old['DateCompleted'] != today_date
        #     ]
        #     ac_pa_client360_autocomplete = pd.concat(
        #         [ac_pa_client360_autocomplete_filtered, combine_pa_autocomplete],
        #         ignore_index=True
        #     )
        # else:
        #     logging.info(f"Creating new file: {autocomplete_path}")
        #     ac_pa_client360_autocomplete = combine_pa_autocomplete.copy()
        # # Save the persistent parquet file
        # ac_pa_client360_autocomplete.to_parquet(autocomplete_path, index=False)
        # logging.info(f"Saved persistent file: {autocomplete_path}")

        # -------------------------------------------------------------------------
        # --- 24. Export Autocomplete to Excel ---
        # -------------------------------------------------------------------------
        logging.info("Step 24: Exporting autocomplete to Excel...")

        excel_path = AC_PATH / "pa_client360_autocomplete.xlsx"

        # --- START FIX ---
        # Use ExcelWriter to set the date format
        with pd.ExcelWriter(
            excel_path,
            engine='openpyxl',
            datetime_format='MM/DD/YYYY H:MM:SS',
            date_format='MM/DD/YYYY'
        ) as writer:
            ac_pa_client360_autocomplete.to_excel(writer, sheet_name="autocomplete", index=False)
        # --- END FIX ---

        logging.info(f"Exported autocomplete Excel: {excel_path}")

        # -------------------------------------------------------------------------
        # --- 25. Create and Export Detail File ---
        # -------------------------------------------------------------------------
        logging.info("Step 25: Creating and exporting detail file...")

        # Based on tmp_pa_C360_4ac_count_pre
        df_detail = tmp_pa_c360_4ac_count_pre.copy()

        # PA_result map (from detail proc sql)
        pa_result_map = {
            'Product Appropriateness assessed outside Client 360': 'Product Appropriate',
            'Not Appropriate - Rationale':                          'Product Not Appropriate',
            'Client declined product appropriateness assessment':   'Client declined',
            'Product Appropriate':                                  'Product Appropriate',
        }

        df_detail['PA_result'] = df_detail['IS_PROD_APRP_FOR_CLNT'].map(pa_result_map).fillna('Missing')

        # Filter
        filter_cond = df_detail['PA_result'].isin([
            'Product Not Appropriate',
            'Missing',
            'Product Appropriateness assessed outside Client 360'
        ])
        df_detail_filtered = df_detail[filter_cond]

        # Select and rename columns
        pa_client360_detail = pd.DataFrame({
            'event_month':              df_detail_filtered['segment10'],
            'reporting_date':           pd.to_datetime(df_detail_filtered['DateCompleted']),
            'event_week_ending':        pd.to_datetime(df_detail_filtered['SnapDate']),
            'event_date':               pd.to_datetime(df_detail_filtered['EVNT_DT']),
            'event_timestamp':          pd.to_datetime(df_detail_filtered['EVNT_TMSTMP']),
            'opportunity_id':           df_detail_filtered['OPPOR_ID'],
            'opportunity_type':         df_detail_filtered['OPPOR_REC_TYP'],
            'product_code':             df_detail_filtered['PROD_CD'],
            'product_category_name':    df_detail_filtered['PROD_CATG_NM'],
            'product_family_name':      df_detail_filtered['ASCT_PROD_FMLY_NM'],
            'product_name':             df_detail_filtered['PROD_SRVC_NM'],
            'oppor_stage_nm':           df_detail_filtered['OPPOR_STAGE_NM'],
            'tool_used':                df_detail_filtered['TOOL_USED'],
            'tool_nm':                  df_detail_filtered['ADVC_TOOL_NM'],
            'PA_result':                df_detail_filtered['PA_result'],
            'PA_rationale':             df_detail_filtered['CLNT_RTNL_TXT'],
            'PA_rationale_validity':    df_detail_filtered['prod_not_aprp_rtnl_txt_cat'],
            'employee_id':              df_detail_filtered['RBC_OPPOR_OWN_ID'],
            'job_code':                 df_detail_filtered['OCCPT_JOB_CD'],
            'position_title':           df_detail_filtered['HR_POSN_TITL_EN'],
            'empolyee_transit':         df_detail_filtered['ORG_UNT_NO'],  # Note: original typo preserved
            'position_start_date':      pd.to_datetime(df_detail_filtered['POSN_STRT_DT']),
        })

        # Export detail file
        detail_excel_path = DATAOUT_PATH / f"pa_client360_detail_{runday}.xlsx"

        # --- START FIX ---
        # Use ExcelWriter to manually apply formats to each date/datetime column
        with pd.ExcelWriter(detail_excel_path, engine='openpyxl') as writer:
            pa_client360_detail.to_excel(writer, sheet_name="detail", index=False)

            # Get the worksheet
            worksheet = writer.sheets['detail']

            # Define formats
            date_format      = 'DD/MM/YYYY'
            timestamp_format = 'M/D/YYYY H:MM:SS AM/PM'

            # Find column letters by header name
            col_letters = {}
            for cell in worksheet[1]:  # Get header row
                col_letters[cell.value] = cell.column_letter

            # Columns to format
            date_cols = ['reporting_date', 'event_week_ending', 'event_date', 'position_start_date']

            # Apply date format
            for col_name in date_cols:
                if col_name in col_letters:
                    col_letter = col_letters[col_name]
                    for cell in worksheet[col_letter][1:]:  # Skip header
                        if cell.value:
                            cell.number_format = date_format

            # Apply timestamp format
            if 'event_timestamp' in col_letters:
                col_letter = col_letters['event_timestamp']
                for cell in worksheet[col_letter][1:]:  # Skip header
                    if cell.value:
                        cell.number_format = timestamp_format
        # --- END FIX ---

        logging.info(f"Exported detail file: {detail_excel_path}")

        # -------------------------------------------------------------------------
        # --- 26. Export Pivot Table ---
        # -------------------------------------------------------------------------
        logging.info("Step 26: Exporting Pivot table...")

        pivot_excel_path = AC_PATH / "pa_client360_Pivot.xlsx"

        ac_pa_client360_autocomplete['DateCompleted'] = pd.to_datetime(
            ac_pa_client360_autocomplete['DateCompleted'], errors='coerce'
        ).dt.date
        ac_pa_client360_autocomplete['SnapDate'] = pd.to_datetime(
            ac_pa_client360_autocomplete['SnapDate'], errors='coerce'
        ).dt.date

        with pd.ExcelWriter(
            pivot_excel_path,
            engine='openpyxl',
            datetime_format='MM/DD/YYYY',
            date_format='MM/DD/YYYY'
        ) as writer:
            ac_pa_client360_autocomplete.to_excel(writer, sheet_name="Autocomplete", index=False)

        logging.info(f"Exported pivot file: {pivot_excel_path}")

        logging.info("=" * 50)
        logging.info("Python script execution completed successfully.")
        logging.info("=" * 50)

    except Exception as e:
        logging.error("=" * 50)
        logging.error(f"SCRIPT FAILED: An unhandled exception occurred: {e}")
        logging.error("=" * 50, exc_info=True)
    finally:
        logging.shutdown()


# =============================================================================
# SQL Server Load
# =============================================================================

def sql_server_load():
    """
    Reads the autocomplete Excel output and bulk-inserts into SQL Server
    staging table STG_REG_AUTOCOMPLETE_C86.
    """
    import openpyxl
    import pyodbc

    logging.info("Starting the script to process client360 Autocomplete data.")
    logging.info("Starting sql server load")

    excel_path = r"sas/RSD/REG_DEV/C86/output/product_appropriateness/client360/pa_client360_autocomplete.xlsx"

    with open(r'sql.json', "r") as f:
        secrets = json.load(f)

    # Define SQL Server connection string
    SQL_SERVER   = secrets["SQL_SERVER"]
    SQL_DATABASE = secrets["SQL_DATABASE"]
    UID          = secrets["SQL_UID"]
    PWD          = secrets["SQL_PWD"]
    SQL_TABLE    = "STG_REG_AUTOCOMPLETE_C86"

    conn_str = (
        f'DRIVER={{ODBC Driver 17 for SQL Server}};'
        f'SERVER={SQL_SERVER};'
        f'DATABASE={SQL_DATABASE};'
        f'UID={UID};'
        f'PWD={PWD};'
        f'TrustServiceCertificate=Yes;'
        f'Connection Timeout=10;'
    )

    # --- Discover expected columns from the target table ---
    conn_probe = pyodbc.connect(conn_str)
    cursor_probe = conn_probe.cursor()
    cursor_probe.execute(f"SELECT TOP 1 * FROM {SQL_DATABASE}.dbo.{SQL_TABLE}")

    expected_columns_raw = [column[0] for column in cursor_probe.description]
    expected_columns     = [column[0] for column in cursor_probe.description]
    print(expected_columns_raw)

    identity_column = 'RowID'
    if identity_column in expected_columns:
        expected_columns.remove(identity_column)
    if identity_column in expected_columns_raw:
        expected_columns_raw.remove(identity_column)

    print(expected_columns_raw)
    cursor_probe.close()
    conn_probe.close()

    # --- Check if the file exists ---
    if not os.path.exists(excel_path):
        print(f"File not found: {excel_path}. Please check the file path.")
        return

    df = pd.read_excel(excel_path)

    # Column rename map to align with SQL Server schema
    column_rename_map = {
        'ControlRisk': 'Control Risk',
        'segment':     'Segment',
        'segment2':    'Segment2',
        'segment3':    'Segment3',
        'segment4':    'Segment4',
        'segment5':    'Segment5',
        'segment6':    'Segment6',
        'segment7':    'Segment7',
        'segment8':    'Segment8',
        'segment9':    'Segment9',
        'segment10':   'Segment10',
        'segment11':   'Segment11',
        'segment12':   'Segment12',
        'segment13':   'Segment13',
        'segment14':   'Segment14',
        'segment15':   'Segment15',
        'segment16':   'Segment16',
        'segment17':   'Segment17',
        'segment18':   'Segment18',
        'segment19':   'Segment19',
        'segment20':   'Segment20',
        'Commentcode': 'CommentCode',
        'Comments.':   'comments',
    }
    df.rename(columns=column_rename_map, inplace=True)

    # Type coercion
    if 'Amount' in df.columns:
        df['Amount'] = pd.to_numeric(df['Amount'], errors='coerce')

    df['Volume']        = pd.to_numeric(df['Volume'],        errors='coerce')
    df['DateCompleted'] = pd.to_datetime(df['DateCompleted'], errors='coerce')
    df['SnapDate']      = pd.to_datetime(df['SnapDate'],      errors='coerce')

    # Ensure all expected columns are present in the DataFrame
    for col in expected_columns:
        if col not in df.columns:
            df[col] = np.nan  # Add missing columns with NaN values

    # Reorder columns to match the expected order
    df = df[expected_columns]

    # Clean up data (replace empty strings or NaN with None)
    df = df.map(lambda x: None if pd.isna(x) or str(x).strip() == "" else x)

    print(df.columns)
    print(len(df.columns))

    # --- Truncate Segment3 to avoid SQL column length overflow ---
    df['Segment3'] = df['Segment3'].astype(str).str.slice(0, 99)
    df['Segment3'] = df['Segment3'].replace({'nan': None, 'NaT': None})

    records = df.where(pd.notnull(df), None).values.tolist()

    # --- Bulk insert ---
    conn   = None
    cursor = None
    try:
        conn   = pyodbc.connect(conn_str)
        cursor = conn.cursor()
        cursor.fast_executemany = True

        placeholders = ', '.join(['?'] * len(expected_columns_raw))
        columns_str  = ', '.join([f'[{col}]' for col in expected_columns_raw])
        insert_sql   = f"INSERT INTO dbo.{SQL_TABLE} ({columns_str}) VALUES ({placeholders})"

        print(placeholders)
        print(columns_str)
        print(insert_sql)

        logging.info(f"Inserting {len(records)} records into {SQL_DATABASE}.dbo.{SQL_TABLE}...")
        logging.info("Performing bulk insert to sql server")
        cursor.executemany(insert_sql, records)
        conn.commit()
        logging.info(f"Inserted {len(records)} records into {SQL_DATABASE}.dbo.{SQL_TABLE}.")
        print("Data inserted successfully.")

    except Exception as e:
        print(f"An error occurred: {e}")
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()
        logging.info("Database connection closed.")


# =============================================================================
# Entry Point
# =============================================================================

if __name__ == "__main__":
    main()
    sql_server_load()
